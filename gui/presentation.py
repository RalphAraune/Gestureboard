# gui/presentation.py

import html
import os
import sys
import time
from collections import deque

import cv2
import fitz

# MediaPipe may fail to load its DLL on some Python setups. Make it optional so
# the rest of the app still starts; feature code guards against mp being None.
try:
    import mediapipe as mp
except Exception:
    print("MediaPipe import error:", sys.exc_info()[1])
    mp = None

# PyAutoGUI is used for the Index Finger presentation cursor. Keep it optional.
try:
    import pyautogui
except Exception:
    pyautogui = None

from PyQt5.QtCore import (
    Qt, QTimer, pyqtSignal, QObject, QPoint, QPointF, QRect, QRectF, QSize
)
from PyQt5.QtGui import (
    QImage, QPixmap, QPainter, QPainterPath, QFont, QColor, QBrush, QPen
)
from PyQt5.QtWidgets import (
    QWidget,
    QLabel,
    QPushButton,
    QComboBox,
    QFileDialog,
    QVBoxLayout,
    QHBoxLayout,
    QGridLayout,
    QFrame,
    QMessageBox,
    QDialog,
    QDialogButtonBox,
    QProgressDialog,
    QSlider,
    QSizePolicy,
    QSplitter,
    QSplitterHandle,
    QStackedWidget,
    QButtonGroup,
)

from core.gesture_detector import GestureDetector, SwipeDetector, WaveDetector
from widgets.annotation_bar import FloatingAnnotationBar
from gui.virtual_mouse import (
    CameraWorker, OneEuroFilter, looks_like_hand, native_move,
)
from core.presentation_controller import PresentationController

try:
    from pptx import Presentation as PptxPresentation
except Exception:
    PptxPresentation = None


# ============================================================
# CAMERA WORKER FOR PRESENTATIONS
#
# The Virtual Mouse's worker passes on a hand only when MediaPipe is
# confident and the hand has been seen for two frames in a row. A
# fast swipe blurs the hand and drops below that confidence, so the
# frames in the middle of a swipe came back empty. This worker also
# keeps the newest hand that has real hand proportions whatever the
# confidence, used only to follow swipes; every other gesture still
# needs the fully checked hand.
# ============================================================

class PresentationCameraWorker(CameraWorker):

    def __init__(self, *args):
        super().__init__(*args)
        self.raw_hand = None

    def _checked_hand(self, results):

        raw = None

        if results is not None and results.multi_hand_landmarks:
            candidate = results.multi_hand_landmarks[0]
            if looks_like_hand(candidate.landmark):
                raw = candidate

        self.raw_hand = raw

        return super()._checked_hand(results)


# ============================================================
# SLIDE INK  (drawings kept per slide)
#
# One store shared by the slide preview and the full screen view,
# so a drawing made in either shows in both. Each slide keeps its
# own ink: going to the next slide shows that slide's (empty) ink,
# and coming back shows the drawing again. Ink is kept in slide
# image pixels, so it stays on what it marked at any window size.
# ============================================================

INK_UNDO_LIMIT = 30

MARKER_ALPHA = 100


class SlideInk(QObject):

    # The ink on screen changed (stroke, undo, slide change, ...).
    changed = pyqtSignal()

    # Drawing was switched on or off.
    enabledChanged = pyqtSignal(bool)

    def __init__(self, parent=None):
        super().__init__(parent)

        # slide index -> [ink image, undo list, redo list]
        self.pages = {}

        self.index = 0
        self.size = QSize()

        self.enabled = False

        self.tool = "pen"              # pen / highlighter / eraser
        self.color = QColor("#E81123")
        self.brush_size = 4            # on-screen pixels

        self._base = None
        self._path = None
        self._last = None

        # Goes up whenever an ink image changes, so the views know when
        # to rescale their cached copy (see SlideView.paintEvent).
        self.version = 0

    # --------------------------------------------------------

    def set_enabled(self, enabled):

        enabled = bool(enabled)

        if enabled != self.enabled:
            self.enabled = enabled
            self.enabledChanged.emit(enabled)

    def set_slide(self, index, size):

        self.end_stroke()
        self.index = index
        self.size = QSize(size)
        self.version += 1
        self.changed.emit()

    def ink(self):

        entry = self.pages.get(self.index)
        return entry[0] if entry else None

    def has_any(self):
        """True when some slide still has something drawn on it."""

        blank = None

        for ink, _undo, _redo in self.pages.values():
            if blank is None or blank.size() != ink.size():
                blank = QImage(ink.size(), ink.format())
                blank.fill(Qt.transparent)
            if ink != blank:
                return True

        return False

    def _entry(self):

        entry = self.pages.get(self.index)

        if entry is None and not self.size.isEmpty():
            ink = QImage(self.size, QImage.Format_ARGB32_Premultiplied)
            ink.fill(Qt.transparent)
            entry = self.pages[self.index] = [ink, [], []]

        return entry

    def stroke_width(self):
        """Width of the current tool on screen, in pixels."""

        if self.tool == "eraser":
            return self.brush_size * 4
        if self.tool == "highlighter":
            return max(10, self.brush_size * 3)
        return self.brush_size

    # --------------------------------------------------------
    # STROKES  (points in slide pixels)
    #
    # A pen or marker stroke is not painted into the ink image while
    # it is drawn: the views paint it on top (live_stroke), and it is
    # painted in once when it ends. Copying the slide-sized image and
    # redrawing the whole stroke on every move was the drawing lag.
    # The eraser works on the image straight away (a short line each
    # move, no copy).
    # --------------------------------------------------------

    def live_stroke(self):
        """(path, pen) of the pen / marker stroke being drawn, or None."""

        if self._path is None or self.tool == "eraser":
            return None

        return self._path, self._stroke_pen()

    def _stroke_pen(self):

        color = QColor(self.color)
        color.setAlpha(MARKER_ALPHA if self.tool == "highlighter" else 235)

        return QPen(color, self._width, Qt.SolidLine, Qt.RoundCap, Qt.RoundJoin)

    def begin_stroke(self, point, width):

        entry = self._entry()

        if entry is None or not self.enabled:
            return

        self._base = entry[0].copy()
        self._width = width
        self._last = QPointF(point)

        # A tiny first segment so a single click leaves a dot.
        self._path = QPainterPath(self._last)
        self._path.lineTo(self._last + QPointF(0.01, 0.01))

        self._paint(self._last)

    def extend_stroke(self, point):

        if self._path is None:
            return

        point = QPointF(point)
        self._path.lineTo(point)
        self._paint(point)

    def end_stroke(self):

        if self._path is None:
            return

        entry = self._entry()

        # Paint the finished pen / marker stroke in, in one go, so a
        # see-through marker stays even where it overlaps itself.
        if self.tool != "eraser":
            painter = QPainter(entry[0])
            painter.setRenderHint(QPainter.Antialiasing)
            painter.setPen(self._stroke_pen())
            painter.setBrush(Qt.NoBrush)
            painter.drawPath(self._path)
            painter.end()

        entry[1].append(self._base)
        if len(entry[1]) > INK_UNDO_LIMIT:
            entry[1].pop(0)
        entry[2].clear()

        self._base = None
        self._path = None
        self.version += 1
        self.changed.emit()

    def _paint(self, point):

        entry = self._entry()

        if self.tool == "eraser":
            # Erase as the cursor moves; the slide itself is never touched.
            painter = QPainter(entry[0])
            painter.setRenderHint(QPainter.Antialiasing)
            painter.setCompositionMode(QPainter.CompositionMode_Clear)
            painter.setPen(QPen(Qt.black, self._width, Qt.SolidLine, Qt.RoundCap, Qt.RoundJoin))
            painter.drawLine(self._last, point)
            painter.end()
            self.version += 1

        self._last = point
        self.changed.emit()

    # --------------------------------------------------------
    # HISTORY
    # --------------------------------------------------------

    def undo(self):

        entry = self.pages.get(self.index)

        if not entry or not entry[1]:
            return

        entry[2].append(entry[0])
        entry[0] = entry[1].pop()
        self.version += 1
        self.changed.emit()

    def redo(self):

        entry = self.pages.get(self.index)

        if not entry or not entry[2]:
            return

        entry[1].append(entry[0])
        entry[0] = entry[2].pop()
        self.version += 1
        self.changed.emit()

    def clear_slide(self):
        """Clear this slide's drawing (undo brings it back)."""

        self.end_stroke()

        entry = self.pages.get(self.index)

        if not entry:
            return

        entry[1].append(entry[0])
        if len(entry[1]) > INK_UNDO_LIMIT:
            entry[1].pop(0)
        entry[2].clear()

        blank = QImage(entry[0].size(), QImage.Format_ARGB32_Premultiplied)
        blank.fill(Qt.transparent)
        entry[0] = blank

        self.version += 1
        self.changed.emit()

    def clear_all(self):
        """Forget the drawings on every slide."""

        self._base = None
        self._path = None
        self.pages = {}
        self.version += 1
        self.changed.emit()


# ============================================================
# SLIDE VIEW  (slide image + its ink)
# ============================================================

class SlideView(QWidget):
    """Shows a slide with its ink from a shared SlideInk.

    The page preview and the full screen view each have one; both
    draw into the same SlideInk while drawing is switched on.
    """

    def __init__(self, ink, background="#0B1A30", empty_text="", parent=None):
        super().__init__(parent)

        self.ink = ink
        self.slide = None              # QPixmap of the current slide
        self._scaled = None            # slide scaled to the current size
        self._ink_cache = None         # (key, ink scaled to the current size)
        self.background = QColor(background)
        self.empty_text = empty_text

        self.setMinimumSize(240, 150)
        self.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)
        self.setMouseTracking(True)

        ink.changed.connect(self.update)
        ink.enabledChanged.connect(self._sync_cursor)
        self._sync_cursor(ink.enabled)

    def _sync_cursor(self, enabled):

        self.setCursor(Qt.CrossCursor if enabled else Qt.ArrowCursor)

    # --------------------------------------------------------

    def set_slide(self, pixmap):

        self.slide = pixmap
        self._scaled = None
        self.update()

    def has_slide(self):
        return self.slide is not None

    def slide_rect(self):
        """The rect the slide is drawn into (centred, keeping its shape)."""

        if self.slide is None:
            return QRect(self.rect())

        size = self.slide.size().scaled(self.size(), Qt.KeepAspectRatio)
        rect = QRect(QPoint(0, 0), size)
        rect.moveCenter(self.rect().center())
        return rect

    def _to_slide(self, pos):

        rect = self.slide_rect()
        scale = self.slide.width() / max(1, rect.width())
        return QPointF((pos.x() - rect.x()) * scale, (pos.y() - rect.y()) * scale), scale

    # --------------------------------------------------------
    # MOUSE  (a real mouse, or the hand cursor pressing it)
    # --------------------------------------------------------

    def mousePressEvent(self, event):

        if event.button() != Qt.LeftButton or self.slide is None:
            return

        if not self.ink.enabled:
            return

        point, scale = self._to_slide(event.pos())
        self.ink.begin_stroke(point, self.ink.stroke_width() * scale)

    def mouseMoveEvent(self, event):

        if not (event.buttons() & Qt.LeftButton) or self.slide is None:
            return

        point, _scale = self._to_slide(event.pos())
        self.ink.extend_stroke(point)

    def mouseReleaseEvent(self, event):

        if event.button() == Qt.LeftButton:
            self.ink.end_stroke()

    # --------------------------------------------------------

    def paintEvent(self, event):

        painter = QPainter(self)
        painter.setRenderHint(QPainter.Antialiasing, True)
        painter.setRenderHint(QPainter.SmoothPixmapTransform, True)

        painter.fillRect(self.rect(), self.background)

        if self.slide is not None:

            rect = self.slide_rect()

            # Scale once per size, not on every stroke repaint.
            if self._scaled is None or self._scaled.size() != rect.size():
                self._scaled = self.slide.scaled(
                    rect.size(), Qt.KeepAspectRatio, Qt.SmoothTransformation
                )

            painter.drawPixmap(rect.topLeft(), self._scaled)

            ink = self.ink.ink()
            if ink is not None and ink.size() == self.slide.size():

                # Rescale the ink only when it changed, not on every
                # repaint while a stroke is drawn.
                key = (self.ink.index, self.ink.version, rect.size())

                if self._ink_cache is None or self._ink_cache[0] != key:
                    self._ink_cache = (key, QPixmap.fromImage(ink.scaled(
                        rect.size(), Qt.IgnoreAspectRatio, Qt.SmoothTransformation
                    )))

                painter.drawPixmap(rect.topLeft(), self._ink_cache[1])

            # The stroke being drawn, on top (slide pixels -> view).
            live = self.ink.live_stroke()
            if live is not None:
                path, pen = live
                scale = rect.width() / max(1, self.slide.width())
                painter.save()
                painter.translate(rect.topLeft())
                painter.scale(scale, scale)
                painter.setPen(pen)
                painter.setBrush(Qt.NoBrush)
                painter.drawPath(path)
                painter.restore()

        elif self.empty_text:

            painter.setPen(QColor("#7890B5"))
            painter.setFont(QFont("Segoe UI", 14))
            painter.drawText(self.rect(), Qt.AlignCenter, self.empty_text)

        painter.end()


# ============================================================
# SMALL WIDGETS USED BY THE PAGE
# ============================================================

class ToneLabel(QLabel):
    """Status label whose colour follows its text.

    `rules` is a list of (substring, tone) pairs checked in order; the
    stylesheet colours the label through its `tone` property.
    """

    def __init__(self, text="", rules=(), parent=None):
        super().__init__(parent)
        self._rules = list(rules)
        self.setText(text)

    def setText(self, text):
        super().setText(text)

        lowered = text.lower()
        tone = next(
            (tone for key, tone in self._rules if key in lowered),
            "idle",
        )

        # Re-polish only on a change: camera text updates every frame.
        if tone != self.property("tone"):
            self.setProperty("tone", tone)
            self.style().unpolish(self)
            self.style().polish(self)


class CameraFeed(QWidget):
    """Camera preview that fills whatever space it gets.

    The frame keeps its aspect ratio and is centred on a rounded well,
    so the panel can be dragged to any size.
    """

    def __init__(self, parent=None):
        super().__init__(parent)

        self._pixmap = None
        self._text = ""
        self._bg = QColor("#DCE7F3")
        self._fg = QColor("#5A6B93")

        self.setMinimumSize(180, 110)
        self.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)

    def sizeHint(self):
        return QSize(320, 240)

    def set_colors(self, background, text):
        self._bg = QColor(background)
        self._fg = QColor(text)
        self.update()

    def setText(self, text):
        self._text = text
        self._pixmap = None
        self.update()

    def setPixmap(self, pixmap):
        self._pixmap = pixmap
        self.update()

    def paintEvent(self, event):

        painter = QPainter(self)
        painter.setRenderHint(QPainter.Antialiasing)
        painter.setRenderHint(QPainter.SmoothPixmapTransform)

        well = QPainterPath()
        well.addRoundedRect(QRectF(self.rect()), 10, 10)
        painter.setClipPath(well)
        painter.fillRect(self.rect(), self._bg)

        if self._pixmap is not None and not self._pixmap.isNull():

            scaled = self._pixmap.size().scaled(
                self.size(), Qt.KeepAspectRatio
            )
            target = QRect(QPoint(0, 0), scaled)
            target.moveCenter(self.rect().center())

            frame = QPainterPath()
            frame.addRoundedRect(QRectF(target), 10, 10)
            painter.setClipPath(frame)
            painter.drawPixmap(target, self._pixmap)

        elif self._text:

            painter.setPen(self._fg)
            painter.setFont(QFont("Segoe UI", 10))
            painter.drawText(self.rect(), Qt.AlignCenter, self._text)

        painter.end()


class _GripHandle(QSplitterHandle):
    """Splitter handle drawn as a small pill so it reads as draggable."""

    def paintEvent(self, event):

        painter = QPainter(self)
        painter.setRenderHint(QPainter.Antialiasing)

        color = QColor(self.splitter().property("gripColor") or "#AAC0E1")

        if not self.underMouse():
            color.setAlpha(150)

        length = 36
        thickness = 4

        if self.orientation() == Qt.Horizontal:
            pill = QRectF(0, 0, thickness, length)
        else:
            pill = QRectF(0, 0, length, thickness)

        pill.moveCenter(QRectF(self.rect()).center())

        painter.setPen(Qt.NoPen)
        painter.setBrush(color)
        painter.drawRoundedRect(pill, 2, 2)
        painter.end()

    def enterEvent(self, event):
        self.update()
        super().enterEvent(event)

    def leaveEvent(self, event):
        self.update()
        super().leaveEvent(event)


class GripSplitter(QSplitter):

    def createHandle(self):
        return _GripHandle(self.orientation(), self)


class PresentationPage(QWidget):

    # Signal can be connected to your MainWindow if needed.
    presentationAction = pyqtSignal(str)

    # Live annotation state changed: "off", "on" or "paused".
    liveAnnotationChanged = pyqtSignal(str)

    @property
    def annotation_mode(self):
        """Drawing is on (the hand gestures draw instead of presenting)."""
        return self.live_annotation == "on"

    def __init__(self, parent=None):
        super().__init__(parent)

        self.parent_window = parent

        # ---------------------------------------------------------
        # PRESENTATION
        # ---------------------------------------------------------

        self.file_path = None
        self.file_type = "Auto Detect"

        self.pdf_document = None
        self.current_page = 0
        self.total_pages = 0
        self.ppt_slides = None

        # In-app fullscreen presentation window (gesture controlled).
        self.fullscreen_window = None

        # ---------------------------------------------------------
        # CAMERA
        # ---------------------------------------------------------

        self.camera = None
        self.camera_index = 0

        # Background camera + hand tracking (see start_camera).
        self.camera_worker = None

        # MediaPipe is optional: if it failed to load, disable hand tracking
        # but keep the page (and the rest of the app) fully functional.
        if mp is not None:
            self.mp_hands = mp.solutions.hands
            self.mp_drawing = mp.solutions.drawing_utils

            # Same fast model as the Virtual Mouse (the heavy model
            # roughly doubled the delay for no better gestures), but
            # lower confidences: a hand partly out of the picture (only
            # the fingers showing) or blurred by a swipe is still found
            # and followed. Gestures other than swipes still need the
            # worker's full hand checks.
            self.hands = self.mp_hands.Hands(
                static_image_mode=False,
                max_num_hands=1,
                model_complexity=0,
                min_detection_confidence=0.5,
                min_tracking_confidence=0.5,
            )
        else:
            self.mp_hands = None
            self.mp_drawing = None
            self.hands = None

        self.gesture_detector = GestureDetector()
        self.swipe_detector = SwipeDetector()

        # A fast swipe loses the hand for a few frames (blur, and the
        # worker re-confirming it); only a longer gap ends the swipe.
        self.swipe_detector.GAP = 0.4
        self.wave_detector = WaveDetector()
        self.presentation_controller = PresentationController()

        self.detected_gesture = "None"
        self.current_action = "Ready"

        # ---------------------------------------------------------
        # ANNOTATION WHILE PRESENTING
        # ---------------------------------------------------------

        # Drawings on the slides, kept per slide and shared by the
        # preview and the full screen view.
        self.slide_ink = SlideInk(self)

        # Started, paused and stopped from the Annotation page (or with
        # the Three Fingers gesture): "off", "on" or "paused". While
        # "on" the floating toolbar shows over the slides and drawing
        # works; "paused" keeps the drawings but hides the toolbar.
        self.live_annotation = "off"

        # True between Pause and Resume, so Start picks up where the
        # presentation left off.
        self.presentation_paused = False

        self.annotation_toolbar = None

        self._annot_drawing = False

        # Steady-hold tracking for one-shot gestures (_handle_hand).
        self._pose = None
        self._pose_since = 0.0

        self.left_click_cooldown = 0

        # Tracks whether this page has been opened at least once. Used to stop
        # gesture handling from firing at startup (splash screen).
        self._page_opened = False

        # Hand-gesture control switch.
        #
        # IMPORTANT state rules:
        #   * Opening Presentation Control does NOT enable gestures.
        #   * Only the Start Presentation button enables them.
        #   * Stop / Exit Full Screen disables them again.
        #   * Switching to any other sidebar page disables them.
        self.gestures_enabled = False

        # Index Finger cursor state (see move_cursor_from_hand): One Euro
        # filter (a little more responsive than the Virtual Mouse's).
        self._cursor_filter_x = OneEuroFilter(min_cutoff=1.0, beta=0.03)
        self._cursor_filter_y = OneEuroFilter(min_cutoff=1.0, beta=0.03)
        self._cursor_last = 0.0

        # Annotation pen (see move_draw_cursor): heavier smoothing than
        # the pointer, so lines are steady.
        self._draw_filter_x = OneEuroFilter(min_cutoff=0.5, beta=0.008)
        self._draw_filter_y = OneEuroFilter(min_cutoff=0.5, beta=0.008)
        self._draw_last = 0.0
        self._draw_hand = None         # hand target last frame (px)
        self._draw_pen = None          # pen before smoothing (px)
        self._draw_string = None       # anti-shake string end (px)

        # Swipe debounce (see _debounce_swipe) and fingertip trails.
        self._swipe_locked_until = 0.0
        self._swipe_return = None
        self._cursor_frozen_until = 0.0
        self._last_swipe = None
        self._tip_trails = {}

        try:
            # Allow the gesture cursor to reach the screen edges/corners.
            pyautogui.FAILSAFE = False
            self._screen_w, self._screen_h = (
                pyautogui.size() if pyautogui else (1920, 1080)
            )
        except Exception:
            self._screen_w, self._screen_h = (1920, 1080)

        # ---------------------------------------------------------
        # BUILD UI
        # ---------------------------------------------------------

        self.build_ui()

        from widgets import page_background

        page_background.clear_layers(self)

        self.apply_theme()

        # ---------------------------------------------------------
        # CAMERA TIMER
        # ---------------------------------------------------------

        self.camera_timer = QTimer(self)
        self.camera_timer.timeout.connect(self.update_camera)

    # ============================================================
    # SHOW EVENT — start the camera only when the page is shown.
    # ============================================================

    def showEvent(self, event):
        super().showEvent(event)
        self._page_opened = True

        # Opening (or returning to) the page must NEVER auto-start the
        # presentation or enable gesture control. The user has to press
        # Start Presentation explicitly.
        self.gestures_enabled = False
        self._refresh_controls()

        if self.camera is None:
            self.start_camera()

    def deactivate_gestures(self):
        """Disable presentation gestures and close any fullscreen view.

        Called when navigating away from Presentation Control (or on Stop /
        Exit Full Screen) so gestures can never open the presentation while
        the user is on another page.
        """

        # Leaving the page while presenting pauses the presentation, so
        # Start resumes on the same slide.
        if self.gestures_enabled:
            self.presentation_paused = True

        self.gestures_enabled = False
        self.end_annotation_draw()
        self.exit_fullscreen()
        self._refresh_controls()

    # ============================================================
    # UI
    # ============================================================

    def build_ui(self):

        self.setObjectName("presentationPage")

        # Width / height breakpoints flip these (see _apply_density).
        self._narrow = False
        self._short = False

        main_layout = QVBoxLayout(self)
        main_layout.setContentsMargins(32, 24, 32, 24)
        main_layout.setSpacing(14)
        self._main_layout = main_layout

        # --------------------------------------------------------
        # HEADER  (title on the left, gesture on/off on the right)
        # --------------------------------------------------------

        header = QHBoxLayout()
        header.setSpacing(12)

        heading = QVBoxLayout()
        heading.setSpacing(2)

        title = QLabel("Presentation Control")
        title.setObjectName("pageTitle")

        self.subtitle = QLabel(
            "Choose your slides, press Start presenting, then use your "
            "hand to move through them."
        )
        self.subtitle.setObjectName("pageSubtitle")
        self.subtitle.setWordWrap(True)

        heading.addWidget(title)
        heading.addWidget(self.subtitle)

        self.status_gesture = ToneLabel(
            "● Gestures off",
            [("on", "good"), ("annotating", "warn"), ("paused", "warn")],
        )
        self.status_gesture.setObjectName("gesturePill")
        self.status_gesture.setToolTip(
            "Hand gestures only work after you press Start presenting."
        )

        header.addLayout(heading, 1)
        header.addWidget(self.status_gesture, 0, Qt.AlignTop)

        main_layout.addLayout(header)

        # --------------------------------------------------------
        # FILE BAR
        # --------------------------------------------------------

        file_bar = QFrame()
        file_bar.setObjectName("fileBar")

        file_layout = QHBoxLayout(file_bar)
        file_layout.setContentsMargins(8, 8, 12, 8)
        file_layout.setSpacing(12)

        self.upload_button = QPushButton("📂  Choose file")
        self.upload_button.setObjectName("uploadButton")
        self.upload_button.setFixedHeight(38)
        self.upload_button.setCursor(Qt.PointingHandCursor)
        self.upload_button.setToolTip("Open a PDF or PowerPoint (.pptx) file")
        self.upload_button.clicked.connect(self.upload_file)

        self.file_status = QLabel(
            "No file yet. Choose a PDF or PowerPoint (.pptx) file to begin."
        )
        self.file_status.setObjectName("fileStatus")
        self.file_status.setProperty("loaded", False)
        self.file_status.setTextFormat(Qt.RichText)
        # Long file names clip instead of widening the page.
        self.file_status.setSizePolicy(
            QSizePolicy.Ignored, QSizePolicy.Preferred
        )

        self.file_type_label = QLabel("Show")
        self.file_type_label.setObjectName("fieldLabel")

        self.file_type_combo = QComboBox()
        self.file_type_combo.setFixedHeight(38)
        self.file_type_combo.setMinimumWidth(130)
        self.file_type_combo.setToolTip(
            "Which files the file picker lists"
        )
        self.file_type_combo.addItems([
            "All supported",
            "PDF",
            "PowerPoint",
        ])

        # Picker controls stay together on the left; the loaded file
        # fills the rest of the bar after a divider.
        divider = QFrame()
        divider.setObjectName("fileDivider")
        divider.setFixedSize(1, 24)

        file_layout.addWidget(self.upload_button)
        file_layout.addWidget(self.file_type_label)
        file_layout.addWidget(self.file_type_combo)
        file_layout.addSpacing(4)
        file_layout.addWidget(divider)
        file_layout.addSpacing(4)
        file_layout.addWidget(self.file_status, 1)

        main_layout.addWidget(file_bar)

        # --------------------------------------------------------
        # BODY  (stage | side column, drag the gap to resize)
        # --------------------------------------------------------

        self.body_splitter = GripSplitter(Qt.Horizontal)
        self.body_splitter.setObjectName("bodySplitter")
        self.body_splitter.setChildrenCollapsible(False)
        self.body_splitter.setHandleWidth(14)

        # ========================================================
        # STAGE  (slide preview + control deck)
        # ========================================================

        stage = QFrame()
        stage.setObjectName("panel")

        stage_layout = QVBoxLayout(stage)
        stage_layout.setContentsMargins(14, 14, 14, 12)
        stage_layout.setSpacing(12)

        preview_holder = QWidget()

        preview_layout = QGridLayout(preview_holder)
        preview_layout.setContentsMargins(0, 0, 0, 0)

        # Slide with its drawings.
        self.slide_view = SlideView(
            self.slide_ink,
            empty_text=(
                "No slides yet\n\n"
                "Press Choose file to open a PDF or PowerPoint."
            ),
        )
        self.preview_label = self.slide_view

        self.slide_counter = QLabel("")
        self.slide_counter.setObjectName("slideCounter")
        self.slide_counter.setAlignment(Qt.AlignCenter)
        self.slide_counter.hide()

        preview_layout.addWidget(self.slide_view, 0, 0)
        preview_layout.addWidget(
            self.slide_counter, 0, 0, Qt.AlignTop | Qt.AlignRight
        )

        # Floating annotation toolbar: over the preview, or over the
        # full screen view while that is open (see _dock_annotation_bar).
        self.annotation_toolbar = FloatingAnnotationBar(self.slide_ink)
        self.annotation_toolbar.attach(self.slide_view)

        stage_layout.addWidget(preview_holder, 1)

        # Control deck, laid out like a presenter remote:
        #   ‹ Previous  |  ▶ Start presenting  |  Next ›     ■ Stop

        self.previous_button = QPushButton()
        self.previous_button.setObjectName("navButton")
        self.previous_button.setToolTip(
            "Go back one slide\nGesture: swipe your open hand left 👈"
        )
        self.previous_button.clicked.connect(self.manual_previous)

        self.start_button = QPushButton()
        self.start_button.setObjectName("startButton")
        self.start_button.clicked.connect(self.manual_start)

        self.next_button = QPushButton()
        self.next_button.setObjectName("navButton")
        self.next_button.setToolTip(
            "Go forward one slide\nGesture: swipe your open hand right 👉"
        )
        self.next_button.clicked.connect(self.manual_next)

        self.stop_button = QPushButton()
        self.stop_button.setObjectName("stopButton")
        self.stop_button.setToolTip(
            "Close full screen and turn hand gestures off"
        )
        self.stop_button.clicked.connect(self.manual_stop)

        # Shown while presenting: leaves full screen and turns gestures
        # off, keeping the slide so Resume carries on from it.
        self.pause_button = QPushButton()
        self.pause_button.setObjectName("pauseButton")
        self.pause_button.setToolTip(
            "Pause on this slide. Press Resume to carry on from here."
        )
        self.pause_button.clicked.connect(self.manual_pause)

        deck = QHBoxLayout()
        deck.setSpacing(8)

        for button, stretch in (
            (self.previous_button, 2),
            (self.start_button, 3),
            (self.next_button, 2),
        ):
            button.setFixedHeight(44)
            button.setMinimumWidth(0)
            button.setCursor(Qt.PointingHandCursor)
            button.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)
            deck.addWidget(button, stretch)

        deck.addSpacing(10)

        self.pause_button.setFixedHeight(44)
        self.pause_button.setMinimumWidth(0)
        self.pause_button.setCursor(Qt.PointingHandCursor)
        self.pause_button.setSizePolicy(
            QSizePolicy.Expanding, QSizePolicy.Fixed
        )
        deck.addWidget(self.pause_button, 1)

        self.stop_button.setFixedHeight(44)
        self.stop_button.setMinimumWidth(0)
        self.stop_button.setCursor(Qt.PointingHandCursor)
        self.stop_button.setSizePolicy(
            QSizePolicy.Expanding, QSizePolicy.Fixed
        )
        deck.addWidget(self.stop_button, 1)

        stage_layout.addLayout(deck)

        # One line that always says what to do next.
        self.hint_label = QLabel()
        self.hint_label.setObjectName("hintLabel")
        self.hint_label.setWordWrap(True)

        stage_layout.addWidget(self.hint_label)

        self.body_splitter.addWidget(stage)

        # ========================================================
        # SIDE COLUMN  (camera | gesture guide, drag to resize)
        # ========================================================

        self.side_splitter = GripSplitter(Qt.Vertical)
        self.side_splitter.setObjectName("sideSplitter")
        self.side_splitter.setChildrenCollapsible(False)
        self.side_splitter.setHandleWidth(14)
        self.side_splitter.setMinimumWidth(290)

        # -------------------- Camera --------------------

        camera_frame = QFrame()
        camera_frame.setObjectName("panel")

        camera_layout = QVBoxLayout(camera_frame)
        camera_layout.setContentsMargins(14, 12, 14, 12)
        camera_layout.setSpacing(10)

        camera_head = QHBoxLayout()
        camera_head.setSpacing(8)

        camera_title = QLabel("Camera")
        camera_title.setObjectName("panelTitle")

        self.camera_status = ToneLabel(
            "● Connecting…",
            [
                ("hand detected", "good"),
                ("not available", "bad"),
                ("connected", "info"),
                ("shared", "info"),
            ],
        )
        self.camera_status.setObjectName("cameraStatus")

        camera_head.addWidget(camera_title)
        camera_head.addStretch()
        camera_head.addWidget(self.camera_status)

        camera_layout.addLayout(camera_head)

        # Scales with the panel; keeps the camera's aspect ratio.
        self.camera_label = CameraFeed()
        self.camera_label.setText("Starting camera…")

        camera_layout.addWidget(self.camera_label, 1)

        readout = QGridLayout()
        readout.setHorizontalSpacing(12)
        readout.setVerticalSpacing(4)
        readout.setColumnStretch(1, 1)

        self.detected_label = QLabel("No gesture")
        self.status_action = QLabel("Ready")

        for row, (name, value) in enumerate((
            ("Seeing", self.detected_label),
            ("Last action", self.status_action),
        )):
            name_label = QLabel(name)
            name_label.setObjectName("readoutName")

            value.setObjectName("readoutValue")
            value.setSizePolicy(QSizePolicy.Ignored, QSizePolicy.Preferred)

            readout.addWidget(name_label, row, 0)
            readout.addWidget(value, row, 1)

        # Hidden on short windows; the hint line under the deck covers it.
        self._action_row = [name_label, self.status_action]

        camera_layout.addLayout(readout)

        # -------------------- Gesture guide --------------------

        guide_frame = QFrame()
        guide_frame.setObjectName("panel")

        guide_layout = QVBoxLayout(guide_frame)
        guide_layout.setContentsMargins(14, 12, 14, 12)
        guide_layout.setSpacing(8)

        self.guide_title = QLabel("Hand gestures")
        self.guide_title.setObjectName("panelTitle")

        guide_layout.addWidget(self.guide_title)

        # Same gesture list the tutorial and Settings show, split into
        # its two groups behind a segmented switch so it never needs
        # to scroll.
        from gui.tutorial import STEPS

        groups = STEPS[0]["groups"]

        segment_row = QHBoxLayout()
        segment_row.setSpacing(0)

        self.guide_stack = QStackedWidget()
        self._guide_grids = []
        self._guide_icons = []
        self.guide_tabs = QButtonGroup(self)
        self.guide_tabs.setExclusive(True)

        for index, (group_title, items) in enumerate(groups):

            tab = QPushButton(group_title)
            tab.setObjectName("segment")
            tab.setCheckable(True)
            tab.setCursor(Qt.PointingHandCursor)
            tab.setFixedHeight(30)
            tab.setProperty(
                "edge",
                "left" if index == 0
                else "right" if index == len(groups) - 1
                else "mid",
            )
            self.guide_tabs.addButton(tab, index)
            segment_row.addWidget(tab, 1)

            page = QWidget()

            rows = QGridLayout(page)
            rows.setContentsMargins(0, 2, 0, 0)
            rows.setHorizontalSpacing(10)
            rows.setVerticalSpacing(5)
            rows.setColumnStretch(1, 1)
            self._guide_grids.append(rows)

            for row, (icon, gesture, action) in enumerate(items):

                icon_label = QLabel(icon)
                icon_label.setObjectName("gestureIcon")
                icon_label.setFixedSize(28, 24)
                icon_label.setAlignment(Qt.AlignCenter)
                self._guide_icons.append(icon_label)

                gesture_label = QLabel(gesture)
                gesture_label.setObjectName("gestureName")
                gesture_label.setWordWrap(True)

                action_label = QLabel(action)
                action_label.setObjectName("gestureAction")
                action_label.setWordWrap(True)
                action_label.setAlignment(
                    Qt.AlignRight | Qt.AlignVCenter
                )

                rows.addWidget(icon_label, row, 0)
                rows.addWidget(gesture_label, row, 1)
                rows.addWidget(action_label, row, 2)

            rows.setRowStretch(len(items), 1)

            self.guide_stack.addWidget(page)

        self.guide_tabs.buttonClicked[int].connect(
            self.guide_stack.setCurrentIndex
        )
        self.guide_tabs.button(0).setChecked(True)

        guide_layout.addLayout(segment_row)
        guide_layout.addWidget(self.guide_stack, 1)

        # Gesture guide on top, camera feed below it.
        self.side_splitter.addWidget(guide_frame)
        self.side_splitter.addWidget(camera_frame)
        self.side_splitter.setStretchFactor(0, 0)
        self.side_splitter.setStretchFactor(1, 1)

        self._guide_frame = guide_frame
        self._side_fitted = False

        self.body_splitter.addWidget(self.side_splitter)
        self.body_splitter.setStretchFactor(0, 1)
        self.body_splitter.setStretchFactor(1, 0)
        self.body_splitter.setSizes([900, 360])

        for splitter in (self.body_splitter, self.side_splitter):
            splitter.handle(1).setToolTip("Drag to resize")

        main_layout.addWidget(self.body_splitter, 1)

        self._refresh_controls()

    # ============================================================
    # RESPONSIVE LAYOUT
    # ============================================================

    def resizeEvent(self, event):

        super().resizeEvent(event)
        self._apply_density()

        # Once the page has its real size, give the gesture guide only
        # the height its rows need and the camera feed the rest.
        if not self._side_fitted and self.isVisible():
            self._side_fitted = True
            QTimer.singleShot(0, self._fit_side_column)

    def _fit_side_column(self):

        frame = self._guide_frame
        guide = frame.minimumSizeHint().height()

        # Wrapped gesture descriptions need more height on a narrow
        # column; measure every tab at the real width.
        layout = frame.layout()
        if layout.hasHeightForWidth():
            guide = max(guide, layout.totalHeightForWidth(frame.width()))
        total = sum(self.side_splitter.sizes())

        self.side_splitter.setSizes([guide, max(1, total - guide)])

    def _apply_density(self):
        """Tighten spacing and shorten labels on small windows so the
        whole page always fits without scrolling."""

        narrow = self.width() < 980
        short = self.height() < 700

        if (narrow, short) == (self._narrow, self._short):
            return

        self._narrow, self._short = narrow, short

        side = 18 if narrow else 32
        top = 14 if short else 24

        self._main_layout.setContentsMargins(side, top, side, top)
        self._main_layout.setSpacing(10 if short else 14)

        self.subtitle.setVisible(not short)
        self.guide_title.setVisible(not short)

        for widget in self._action_row:
            widget.setVisible(not short)

        for grid in self._guide_grids:
            grid.setVerticalSpacing(1 if short else 5)

        for icon in self._guide_icons:
            icon.setFixedSize(28, 20 if short else 24)

        self.camera_label.setMinimumHeight(80 if short else 110)

        if self._side_fitted:
            QTimer.singleShot(0, self._fit_side_column)
        self.file_type_label.setVisible(not narrow)
        self.side_splitter.setMinimumWidth(250 if narrow else 290)

        self._refresh_controls()

    # ============================================================
    # CONTROL STATE  (buttons, hint line and gesture pill)
    # ============================================================

    def _has_slides(self):

        return bool(self.pdf_document) or bool(
            getattr(self, "ppt_slides", None)
        )

    def _refresh_controls(self):

        if not hasattr(self, "start_button"):
            return

        loaded = self._has_slides()
        presenting = self.gestures_enabled
        narrow = self._narrow

        paused = self.presentation_paused and not presenting

        self.previous_button.setText("‹  Prev" if narrow else "‹  Previous")
        self.next_button.setText("Next  ›")
        self.stop_button.setText("■  Stop")
        self.pause_button.setText("❚❚  Pause")
        self.pause_button.setVisible(presenting)

        if paused:
            slide = self.current_page + 1
            self.start_button.setText(
                "▶  Resume" if narrow else f"▶  Resume from slide {slide}"
            )
            self.start_button.setToolTip(
                f"Carry on presenting from slide {slide}"
            )
        elif presenting:
            self.start_button.setText(
                "⤢  Full screen" if narrow else "⤢  Back to full screen"
            )
            self.start_button.setToolTip(
                "Reopen the slides full screen\nGesture: thumbs up 👍"
            )
        else:
            self.start_button.setText(
                "▶  Start" if narrow else "▶  Start presenting"
            )
            self.start_button.setToolTip(
                "Show the slides full screen and turn hand gestures on\n"
                "Gesture: rock sign 🤘"
            )

        self.start_button.setEnabled(loaded)
        self.previous_button.setEnabled(loaded)
        self.next_button.setEnabled(loaded)
        self.stop_button.setEnabled(presenting or paused)

        if not loaded:
            hint = "Choose a file above. Your slides show up here."
        elif self.annotation_mode and presenting:
            hint = (
                "Drawing is on: hold index and middle fingers together to "
                "draw, swipe your open hand left or right to change "
                "slides, wave your open hand to clear this slide, raise "
                "your pinky 🤙 to close drawing."
            )
        elif self.annotation_mode:
            hint = (
                "Drawing is on: draw on the slide with the mouse, or press "
                "Resume and use your hand. Each slide keeps its own drawing."
            )
        elif paused:
            hint = (
                f"Paused on slide {self.current_page + 1}. Press Resume to "
                "carry on. To draw on your slides, press Start on the "
                "Annotation page."
            )
        elif presenting:
            hint = (
                "Hand gestures are on. Swipe your open hand right for "
                "the next slide, left to go back."
            )
        else:
            hint = (
                "Hand gestures are off. Press Start presenting or show the "
                "rock sign 🤘 to start, or use Previous and Next."
            )

        self.hint_label.setText(hint)

        if self.annotation_mode:
            self.status_gesture.setText("● Annotating")
        elif presenting:
            self.status_gesture.setText("● Gestures on")
        elif paused:
            self.status_gesture.setText("● Paused")
        else:
            self.status_gesture.setText("● Gestures off")

    def _update_counter(self):

        if self._has_slides() and self.total_pages:
            self.slide_counter.setText(
                f"Slide {self.current_page + 1} of {self.total_pages}"
            )
            self.slide_counter.show()
        else:
            self.slide_counter.hide()

    # ============================================================
    # THEME
    #
    # Every colour on this page lives in this one stylesheet, so
    # MainWindow can re-apply it when the user switches light / dark.
    # ============================================================

    def apply_theme(self):

        from settings import theme as app_theme
        from widgets import page_background

        c = self._page_colors = page_background.page_colors()

        dark = app_theme.is_dark()
        field = c["input_bg"] if dark else "#F5F9FD"
        well = c["bg"] if dark else "#DCE7F3"
        focus = c["accent"] if dark else c["sidebar_hover"]

        if dark:
            ok_bg, ok_text, ok_border = "#173B2E", "#5FD39A", "#245A45"
            stop_bg, stop_text, stop_border = "#3A2230", "#F59A93", "#5A3040"
            warn_bg, warn_text = "#3A3220", "#F2C46B"
            info_text = "#8DB4F5"
            disabled_bg, disabled_text = "#132742", "#5A7396"
        else:
            ok_bg, ok_text, ok_border = "#E3F4EA", "#23804F", "#B7E1C7"
            stop_bg, stop_text, stop_border = "#F8E3E1", "#B83227", "#EBC3BF"
            warn_bg, warn_text = "#FBF0D9", "#8A5A00"
            info_text = "#1B4499"
            disabled_bg, disabled_text = "#E4ECF6", "#9AAACA"

        self.camera_label.set_colors(well, c["text_secondary"])

        for splitter in (self.body_splitter, self.side_splitter):
            splitter.setProperty("gripColor", c["input_border"])
            splitter.handle(1).update()

        self.setStyleSheet(
            page_background.LAYER_RULE
            + page_background.combo_rules(c, field)
            + f"""
            QWidget#presentationPage {{
                background: transparent;
            }}

            QLabel {{
                background: transparent;
                color: {c['text']};
            }}

            QLabel#pageTitle {{
                font-size: 28px;
                font-weight: 800;
            }}

            QLabel#pageSubtitle {{
                font-size: 14px;
                color: {c['text_secondary']};
            }}

            QFrame#panel, QFrame#fileBar {{
                background: {c['card']};
                border: 1px solid {c['card_border']};
                border-radius: 14px;
            }}

            QLabel#panelTitle {{
                font-size: 15px;
                font-weight: 700;
            }}

            QFrame#fileDivider {{
                background: {c['card_border']};
                border: none;
            }}

            QLabel#fieldLabel {{
                color: {c['text_secondary']};
                font-size: 12px;
            }}

            /* ---------- Gesture on/off pill ---------- */

            QLabel#gesturePill {{
                background: {c['card']};
                border: 1px solid {c['card_border']};
                color: {c['text_secondary']};
                border-radius: 15px;
                padding: 6px 14px;
                font-size: 13px;
                font-weight: 700;
            }}

            QLabel#gesturePill[tone="good"] {{
                background: {ok_bg};
                border-color: {ok_border};
                color: {ok_text};
            }}

            QLabel#gesturePill[tone="warn"] {{
                background: {warn_bg};
                border-color: {warn_text};
                color: {warn_text};
            }}

            /* ---------- Buttons ---------- */

            QPushButton {{
                background: {field};
                color: {c['text']};
                border: 1px solid {c['input_border']};
                border-radius: 10px;
                font-size: 13px;
                font-weight: 600;
                padding: 6px 14px;
            }}

            QPushButton:hover {{
                background: {c['button_hover']};
            }}

            QPushButton:focus {{
                outline: none;
                border: 2px solid {focus};
            }}

            QPushButton:disabled {{
                background: {disabled_bg};
                color: {disabled_text};
                border: 1px solid {c['card_border']};
            }}

            QPushButton#uploadButton {{
                background: {c['primary']};
                color: {c['primary_text']};
                border: none;
                font-weight: 700;
                padding: 6px 18px;
            }}

            QPushButton#uploadButton:hover {{
                background: {c['sidebar_hover']};
                color: {c['sidebar_text']};
            }}

            QPushButton#navButton, QPushButton#pauseButton {{
                font-size: 14px;
            }}

            QPushButton#pauseButton {{
                font-weight: 600;
            }}

            QPushButton#startButton {{
                background: {c['primary']};
                color: {c['primary_text']};
                border: none;
                font-size: 14px;
                font-weight: 700;
            }}

            QPushButton#startButton:hover {{
                background: {c['sidebar_hover']};
                color: {c['sidebar_text']};
            }}

            QPushButton#startButton:focus {{
                border: 2px solid {c['accent']};
            }}

            QPushButton#startButton:disabled {{
                background: {disabled_bg};
                color: {disabled_text};
            }}

            QPushButton#stopButton {{
                background: {stop_bg};
                color: {stop_text};
                border: 1px solid {stop_border};
            }}

            QPushButton#stopButton:hover {{
                border: 1px solid {stop_text};
            }}

            QPushButton#stopButton:disabled {{
                background: transparent;
                color: {disabled_text};
                border: 1px dashed {c['card_border']};
            }}

            /* ---------- File status ---------- */

            QLabel#fileStatus {{
                color: {c['text_secondary']};
                font-size: 13px;
            }}

            QLabel#fileStatus[loaded="true"] {{
                color: {c['text']};
            }}

            /* ---------- Stage ---------- */

            QLabel#slideCounter {{
                background: rgba(11, 26, 48, 0.82);
                color: #FFFFFF;
                padding: 5px 12px;
                margin: 10px;
                border-radius: 12px;
                font-weight: 600;
                font-size: 12px;
            }}

            QLabel#hintLabel {{
                color: {c['text_secondary']};
                font-size: 12px;
                padding: 0 2px;
            }}

            /* ---------- Camera ---------- */

            QLabel#cameraStatus {{
                color: {c['text_secondary']};
                font-size: 12px;
                font-weight: 600;
            }}

            QLabel#cameraStatus[tone="good"] {{
                color: {ok_text};
            }}

            QLabel#cameraStatus[tone="info"] {{
                color: {info_text};
            }}

            QLabel#cameraStatus[tone="bad"] {{
                color: {stop_text};
            }}

            QLabel#readoutName {{
                color: {c['text_secondary']};
                font-size: 12px;
            }}

            QLabel#readoutValue {{
                font-size: 13px;
                font-weight: 700;
            }}

            /* ---------- Gesture guide ---------- */

            QPushButton#segment {{
                background: {field};
                color: {c['text_secondary']};
                border: 1px solid {c['input_border']};
                border-radius: 0;
                font-size: 12px;
                padding: 4px 8px;
            }}

            QPushButton#segment[edge="left"] {{
                border-top-left-radius: 8px;
                border-bottom-left-radius: 8px;
                border-right: none;
            }}

            QPushButton#segment[edge="right"] {{
                border-top-right-radius: 8px;
                border-bottom-right-radius: 8px;
            }}

            QPushButton#segment:checked {{
                background: {c['primary']};
                color: {c['primary_text']};
                border-color: {c['primary']};
                font-weight: 700;
            }}

            QLabel#gestureIcon {{
                background: {c['button_hover']};
                border-radius: 7px;
                font-family: "Segoe UI Emoji";
                font-size: 14px;
            }}

            QLabel#gestureName {{
                font-size: 12px;
                font-weight: 700;
            }}

            QLabel#gestureAction {{
                font-size: 12px;
                color: {c['text_secondary']};
            }}

            QSplitter::handle {{
                background: transparent;
            }}
            """
        )

        self.update()

    def _set_file_loaded(self, loaded):

        self.file_status.setProperty("loaded", bool(loaded))
        self.file_status.style().unpolish(self.file_status)
        self.file_status.style().polish(self.file_status)

    def _file_summary(self, file_path, kind):

        count = self.total_pages
        slides = "slide" if count == 1 else "slides"
        name = html.escape(os.path.basename(file_path))

        return f"<b>{name}</b>&nbsp;&nbsp;&nbsp;{kind}, {count} {slides}"

    # ============================================================
    # BACKGROUND
    # ============================================================

    def paintEvent(self, event):

        from widgets import page_background

        page_background.paint_page_background(self, self._page_colors)

    # ============================================================
    # STATUS ROW
    # ============================================================

    def add_status_row(self, layout, row, title, value):

        title_label = QLabel(title)
        title_label.setObjectName("statusName")

        value.setObjectName("statusValue")

        layout.addWidget(title_label, row, 0)
        layout.addWidget(
            value,
            row,
            1,
            alignment=Qt.AlignRight
        )

    # ============================================================
    # UPLOAD FILE
    # ============================================================

    def upload_file(self):

        selected_type = self.file_type_combo.currentText()

        if selected_type == "PDF":
            filter_text = "PDF Files (*.pdf)"

        elif selected_type == "PowerPoint":
            filter_text = (
                "PowerPoint Files (*.ppt *.pptx)"
            )

        else:
            filter_text = (
                "Presentation Files "
                "(*.pdf *.ppt *.pptx)"
            )

        file_path, _ = QFileDialog.getOpenFileName(
            self,
            "Select Presentation",
            "",
            filter_text
        )

        if not file_path:
            return

        self.file_path = file_path

        self.presentation_controller.set_file(
            file_path
        )

        extension = (
            os.path.splitext(file_path)[1]
            .lower()
        )

        if extension == ".pdf":
            self.load_pdf(file_path)

        elif extension == ".pptx":
            self.load_powerpoint(file_path)

        elif extension == ".ppt":
            # Legacy binary .ppt files cannot be rendered in-app (python-pptx
            # only reads the modern .pptx format). Give a clear message
            # instead of the cryptic "Package not found" error.
            QMessageBox.warning(
                self,
                "Unsupported PowerPoint Format",
                "Legacy .ppt files are not supported.\n\n"
                "Please open the file in PowerPoint and use "
                "\"Save As\" → PowerPoint Presentation (*.pptx), "
                "then upload the .pptx file."
            )

        else:
            QMessageBox.warning(
                self,
                "Unsupported File",
                "Please select a PDF or PowerPoint (.pptx) file."
            )

    # ============================================================
    # PDF
    # ============================================================

    def load_pdf(self, file_path):

        try:

            if self.pdf_document:
                self.pdf_document.close()

            self.pdf_document = fitz.open(
                file_path
            )

            # A new file starts without drawings.
            self.slide_ink.clear_all()

            self.total_pages = len(
                self.pdf_document
            )

            self.current_page = 0

            self.file_status.setText(
                self._file_summary(file_path, "PDF")
            )

            self._set_file_loaded(True)

            self.render_pdf_page()

            self.update_status()

        except Exception as e:

            QMessageBox.critical(
                self,
                "PDF Error",
                f"Unable to load PDF:\n\n{e}"
            )

    # ============================================================
    # POWERPOINT
    # ============================================================

    def load_powerpoint(self, file_path):

        if PptxPresentation is None:
            QMessageBox.warning(
                self,
                "PowerPoint Unavailable",
                "python-pptx is not installed. "
                "Please install it to view PowerPoint files in-app."
            )
            return

        try:
            prs = PptxPresentation(file_path)

            # A new file starts without drawings.
            self.slide_ink.clear_all()

            slides = [slide for slide in prs.slides]
            self.total_pages = len(slides)
            self.current_page = 0

            self.ppt_slides = [
                self.render_pptx_slide(slide, prs)
                for slide in slides
            ]

            self.file_status.setText(
                self._file_summary(file_path, "PowerPoint")
            )

            self._set_file_loaded(True)

            self.render_pdf_page()

            self.update_status()

        except Exception as e:

            QMessageBox.critical(
                self,
                "PowerPoint Error",
                f"Unable to load PowerPoint:\n\n{e}"
            )

    # ============================================================
    # PPTX RENDER (in-app)
    # ============================================================

    def render_pptx_slide(self, slide, prs):
        """Render a python-pptx slide into a QImage using QPainter."""

        # Use a consistent canvas size (16:9-ish) scaled for display.
        width = 1600
        height = 900

        img = QImage(
            width,
            height,
            QImage.Format_ARGB32
        )
        img.fill(QColor("#FFFFFF"))

        painter = QPainter(img)
        painter.setRenderHint(
            QPainter.Antialiasing,
            True
        )

        # Slide dimensions from the presentation (EMU -> px).
        try:
            slide_w = prs.slide_width
            slide_h = prs.slide_height
            scale = width / float(slide_w)
        except Exception:
            scale = 1.0

        # Draw text shapes (title & body placeholders / text boxes).
        for shape in slide.shapes:
            try:
                if not shape.has_text_frame:
                    continue
                text = shape.text_frame.text.strip()
                if not text:
                    continue
                if text == "Click to add title":
                    continue

                left = int(shape.left * scale)
                top = int(shape.top * scale)
                s_width = int(shape.width * scale)
                s_height = int(shape.height * scale)

                is_title = (
                    shape.name.lower().startswith(
                        ("title", "subtitle")
                    )
                )

                # Auto-fit font so the full text is visible without clipping.
                body_font = 24
                font_size = (
                    46
                    if is_title
                    else body_font
                )

                box_w = max(30, s_width)
                box_h = max(30, s_height)

                while font_size > 8:
                    painter.setFont(
                        QFont("Segoe UI", font_size)
                    )
                    metrics = painter.fontMetrics()
                    needed_w = metrics.horizontalAdvance(text)
                    # Rough estimate: add padding for wrapped lines.
                    est_lines = max(
                        1,
                        needed_w // max(1, box_w) + 1
                    )
                    total_h = est_lines * metrics.height()
                    if total_h <= box_h or font_size <= 10:
                        break
                    font_size -= 2

                painter.setPen(
                    QColor("#132B5C")
                )
                painter.setFont(
                    QFont("Segoe UI", font_size)
                )

                rect = (
                    left,
                    top,
                    box_w,
                    box_h
                )

                painter.drawText(
                    rect[0],
                    rect[1],
                    rect[2],
                    rect[3],
                    Qt.AlignTop
                    | Qt.AlignLeft
                    | Qt.TextWordWrap,
                    text,
                )

            except Exception:
                continue

        painter.end()

        return img

    # ============================================================
    # PDF RENDER
    # ============================================================

    def render_pdf_page(self):

        # PowerPoint slides render from the in-app image cache.
        if getattr(self, "ppt_slides", None):
            if self.current_page < 0:
                self.current_page = 0
            if self.current_page >= self.total_pages:
                self.current_page = self.total_pages - 1

            img = self.ppt_slides[self.current_page]

            pixmap = QPixmap.fromImage(img)

            # Each slide keeps its own drawing.
            self.slide_view.set_slide(pixmap)
            self.slide_ink.set_slide(self.current_page, pixmap.size())

            self._update_counter()

            self.refresh_fullscreen()
            return

        if not self.pdf_document:
            return

        if self.current_page < 0:
            self.current_page = 0

        if self.current_page >= self.total_pages:
            self.current_page = (
                self.total_pages - 1
            )

        page = self.pdf_document.load_page(
            self.current_page
        )

        matrix = fitz.Matrix(
            1.4,
            1.4
        )

        pix = page.get_pixmap(
            matrix=matrix,
            alpha=False
        )

        image = QImage(
            pix.samples,
            pix.width,
            pix.height,
            pix.stride,
            QImage.Format_RGB888
        )

        pixmap = QPixmap.fromImage(
            image
        )

        # Each slide keeps its own drawing.
        self.slide_view.set_slide(
            pixmap
        )

        self.slide_ink.set_slide(self.current_page, pixmap.size())

        self._update_counter()

        self.refresh_fullscreen()

    # ============================================================
    # PDF THUMBNAILS
    # ============================================================

    def clear_thumbnails(self):

        # Thumbnail strip was removed from the layout; keep this a safe no-op.
        if not hasattr(self, "thumbnail_layout"):
            return

        while self.thumbnail_layout.count():

            item = self.thumbnail_layout.takeAt(0)

            widget = item.widget()

            if widget:
                widget.deleteLater()

    def create_pdf_thumbnails(self):

        # Thumbnail strip was removed from the layout; keep this a safe no-op.
        if not hasattr(self, "thumbnail_layout"):
            return

        self.clear_thumbnails()

        if not self.pdf_document:
            return

        max_thumbnails = min(
            self.total_pages,
            30
        )

        for index in range(max_thumbnails):

            page = self.pdf_document.load_page(
                index
            )

            pix = page.get_pixmap(
                matrix=fitz.Matrix(
                    0.35,
                    0.35
                ),
                alpha=False
            )

            image = QImage(
                pix.samples,
                pix.width,
                pix.height,
                pix.stride,
                QImage.Format_RGB888
            )

            thumbnail = QPixmap.fromImage(
                image
            ).scaled(
                110,
                70,
                Qt.KeepAspectRatio,
                Qt.SmoothTransformation
            )

            button = QPushButton()

            button.setFixedSize(
                120,
                82
            )

            button.setIcon(
                thumbnail
            )

            button.setIconSize(
                thumbnail.size()
            )

            button.clicked.connect(
                lambda checked=False,
                page_index=index:
                self.select_page(page_index)
            )

            self.thumbnail_layout.addWidget(
                button
            )

    def select_page(self, page_index):

        if not self.pdf_document:
            return

        self.current_page = page_index

        self.render_pdf_page()
        self.update_status()

    # ============================================================
    # CAMERA
    # ============================================================

    def _running_vmouse(self):
        """Return the Virtual Mouse page if it is currently tracking."""

        vmouse = getattr(
            self.parent_window, "virtual_mouse", None
        )

        if vmouse is not None and getattr(
            vmouse, "camera_running", False
        ):
            return vmouse

        return None

    def start_camera(self):

        # Do not open a second handle to the same webcam.
        if self.camera is not None:
            return

        # The Virtual Mouse already owns the webcam: reuse its frames and
        # hand landmarks instead of opening the device a second time, so
        # both features keep working together.
        if self._running_vmouse() is not None:

            self._last_shared_frame = None

            # Cheap check for a new Virtual Mouse frame; frequent so a
            # new frame is picked up almost as soon as it exists.
            if not self.camera_timer.isActive():
                self.camera_timer.start(10)

            self.camera_status.setText(
                "● Shared with Virtual Mouse"
            )

            return

        backend = cv2.CAP_DSHOW if hasattr(cv2, "CAP_DSHOW") else cv2.CAP_ANY

        self.camera = cv2.VideoCapture(
            self.camera_index,
            backend
        )

        if not self.camera.isOpened():

            self.camera.release()

            self.camera = cv2.VideoCapture(
                self.camera_index
            )

        if self.camera.isOpened():

            # Same settings as the Virtual Mouse: MJPG for full frame
            # rate, and a one-frame buffer so stale frames never queue up
            # behind the hand (the delay). Unsupported settings are
            # ignored by the camera.
            for prop, value in (
                (cv2.CAP_PROP_FOURCC, cv2.VideoWriter_fourcc(*"MJPG")),
                (cv2.CAP_PROP_FRAME_WIDTH, 640),
                (cv2.CAP_PROP_FRAME_HEIGHT, 480),
                (cv2.CAP_PROP_FPS, 30),
                (cv2.CAP_PROP_BUFFERSIZE, 1),
            ):
                try:
                    self.camera.set(prop, value)
                except Exception:
                    pass

            # Reading the camera and tracking the hand run on background
            # threads (the Virtual Mouse's camera worker), keeping only
            # the newest frame, so the window never waits on them.
            self.camera_worker = PresentationCameraWorker(
                self.camera,
                self.hands,
                self.mp_hands,
                self.mp_drawing,
            )
            self.camera_worker.ready.connect(self._on_worker_frame)
            self.camera_worker.start()

            self.camera_status.setText(
                "● Connected"
            )

        else:

            # Release the failed handle and clear it so a later
            # showEvent()/navigation can try again.
            self.camera.release()
            self.camera = None

            self.camera_status.setText(
                "● Camera not available"
            )

            self.camera_label.setText(
                "No camera found.\nPick one in Camera Setup."
            )

    def _on_worker_frame(self):
        """A new frame from the camera worker (UI thread)."""

        worker = self.camera_worker

        if worker is None:
            return

        result = worker.take()

        if result is None:
            return

        frame, hand, _stamp = result

        # The worker has mirrored the frame, tracked the hand (only
        # real, confirmed hands) and drawn it. raw_hand may be a
        # blurred hand it did not confirm: only swipes use it.
        self.process_frame(frame, hand, raw_hand=worker.raw_hand)

    # ============================================================
    # CAMERA SELECTION
    # ============================================================

    def setActiveCamera(self, index, name):

        if index is None or index == self.camera_index:
            return

        # Open on its own, or mirroring the Virtual Mouse (timer only).
        was_open = (
            self.camera is not None
            or self.camera_timer.isActive()
        )

        self.release_camera()

        self.camera_index = index

        # Reopen on the new device only if this page was using the camera.
        if was_open and self.isVisible():
            self.start_camera()

    def updateCameraName(self, name):

        self.camera_name = name

    # ============================================================
    # RELEASE CAMERA
    #
    # Called by MainWindow when navigating away from this page so the
    # single webcam is freed for other pages (e.g. Virtual Mouse). This
    # prevents two pages from fighting over the same camera device.
    # ============================================================

    def release_camera(self):

        if self.camera_timer.isActive():
            self.camera_timer.stop()

        worker = getattr(self, "camera_worker", None)

        if worker is not None:
            self.camera_worker = None
            worker.stop()

        if self.camera is not None:
            self.camera.release()
            self.camera = None

        # A hand that vanishes with the camera must not keep drawing.
        self._handle_no_hand()

    # ============================================================
    # SHARED CAMERA (from the Virtual Mouse)
    # ============================================================

    def update_camera(self):
        """Timer tick: only used while sharing the Virtual Mouse camera
        (an own camera delivers frames through the camera worker)."""

        if self.camera is None:
            self.update_shared_camera()

    def update_shared_camera(self):

        vmouse = self._running_vmouse()

        if vmouse is None:

            # The Virtual Mouse stopped and freed the webcam: take it
            # over so presentation gestures keep working.
            self.camera_timer.stop()

            if self.isVisible():
                self.start_camera()

            return

        frame = vmouse.get_latest_frame()

        # Only handle each Virtual Mouse frame once, so a gesture is
        # not triggered twice from the same image.
        if frame is None or frame is getattr(
            self, "_last_shared_frame", None
        ):
            return

        self._last_shared_frame = frame

        # The Virtual Mouse already ran MediaPipe, with the same real
        # hand checks, and drew the landmarks; reuse its result.
        self.process_frame(
            frame.copy(),
            vmouse.get_latest_hand(),
            shared=True,
        )

    # ============================================================
    # PROCESS FRAME (gestures + preview)
    # ============================================================

    # Landmarks this close to the frame edge count as cut off.
    EDGE_MARGIN = 0.01

    def _whole_hand(self, hand):
        """True when all 21 points of the hand are inside the picture.

        MediaPipe guesses the points of a hand that is partly out of
        frame; a guessed finger can look like any gesture, so a cut-off
        hand does nothing.
        """

        low, high = self.EDGE_MARGIN, 1.0 - self.EDGE_MARGIN

        try:
            return all(
                low <= point.x <= high and low <= point.y <= high
                for point in hand.landmark
            )
        except Exception:
            return False

    # Points a swipe needs in the picture: the four knuckles it
    # follows (see SwipeDetector.palm_centre). The wrist and fingertips
    # may be out of frame.
    SWIPE_POINTS = (5, 9, 13, 17)

    def _swipe_points_inside(self, hand):

        try:
            return all(
                0.0 <= hand.landmark[i].x <= 1.0
                and 0.0 <= hand.landmark[i].y <= 1.0
                for i in self.SWIPE_POINTS
            )
        except Exception:
            return False

    def process_frame(self, frame, hand, shared=False, raw_hand=None):
        """One camera frame: `hand` is a real, confirmed hand or None.

        `raw_hand` is a hand the camera worker saw but did not confirm
        (usually blurred by fast movement); it is only used to follow
        a swipe.
        """

        gesture = "None"
        cut_off = False

        # Swipes come first and are followed on any frame where the
        # knuckles can be seen: mid-swipe the hand blurs (not confirmed)
        # or reaches the edge of the picture (not whole). Frames without
        # a hand still go in, so a swipe waiting to be confirmed fires.
        swipe_hand = hand if hand is not None else raw_hand

        self._update_tip_trails(swipe_hand, time.monotonic())

        if swipe_hand is not None and not self._swipe_points_inside(swipe_hand):
            swipe_hand = None

        swiped = self._check_swipe(swipe_hand)

        if hand is not None and not self._whole_hand(hand):
            hand = None
            cut_off = True

        if swiped:

            # Stop any drawing stroke the swipe may have started.
            self.end_annotation_draw()
            self._reset_clear_hold()
            self.detected_gesture = "Swipe"
            self._pose = None

            self.camera_status.setText("● Hand detected")

        elif hand is not None:

            gesture = self._detect(hand)
            self._handle_hand(gesture, hand)

            self.camera_status.setText("● Hand detected")

        else:

            self._handle_no_hand()

            if cut_off:
                self.camera_status.setText("● Show your whole hand")
            elif shared:
                self.camera_status.setText("● Shared with Virtual Mouse")
            else:
                self.camera_status.setText("● Connected")

        if cut_off:
            seen = "Hand cut off: show all of it"
        elif gesture in ("None", "Unknown"):
            seen = "No gesture"
        elif self.gestures_enabled:
            seen = gesture
        else:
            seen = gesture + " (gestures off)"

        # A swipe shows its own name for a moment.
        if time.monotonic() > getattr(self, "_swipe_shown_until", 0.0):
            self.detected_label.setText(seen)

        # --------------------------------------------------------
        # CAMERA IMAGE
        # --------------------------------------------------------

        frame_rgb = cv2.cvtColor(
            frame,
            cv2.COLOR_BGR2RGB
        )

        height, width, channel = (
            frame_rgb.shape
        )

        image = QImage(
            frame_rgb.data,
            width,
            height,
            channel * width,
            QImage.Format_RGB888
        )

        self._draw_swipe_trail(frame_rgb)

        # CameraFeed scales it to the panel size when painting.
        self.camera_label.setPixmap(
            QPixmap.fromImage(image)
        )

    # ============================================================
    # ANNOTATION WHILE PRESENTING  (Start / Pause / Stop)
    #
    # Controlled from the Annotation page. While "on", the floating
    # toolbar shows over the slides and the hand gestures draw.
    # ============================================================

    def start_live_annotation(self):

        if not self._has_slides():
            return False

        self._set_live_annotation("on")
        return True

    def pause_live_annotation(self):

        if self.live_annotation == "on":
            self._set_live_annotation("paused")

    def stop_live_annotation(self):
        """End annotating and forget the drawings on every slide."""

        self.slide_ink.clear_all()
        self._set_live_annotation("off")

    def _set_live_annotation(self, state):

        self.live_annotation = state

        self.end_annotation_draw()
        self._reset_clear_hold()

        self.slide_ink.set_enabled(state == "on")

        self._dock_annotation_bar()

        # Show the drawing gestures while annotating.
        if hasattr(self, "guide_tabs"):
            tab = self.guide_tabs.button(1 if state == "on" else 0)
            if tab is not None:
                tab.click()

        self._set_action({
            "on": "Annotation on",
            "paused": "Annotation paused",
            "off": "Annotation stopped",
        }[state])

        self.update_status()
        self.liveAnnotationChanged.emit(state)

    def _dock_annotation_bar(self):
        """Float the toolbar over the full screen view if it is open,
        otherwise over the slide preview."""

        bar = self.annotation_toolbar

        if bar is None:
            return

        viewer = self.fullscreen_window
        target = viewer.slide_view if viewer is not None else self.slide_view

        if bar.parentWidget() is not target:
            bar.attach(target)

        bar.setVisible(self.live_annotation == "on")
        bar.raise_()

    def _reset_clear_hold(self):
        """Forget a half-done wave and empty the circle's red fill."""

        self.wave_detector.reset()

        if self.annotation_toolbar is not None:
            self.annotation_toolbar.set_clear_progress(0)

    def _check_wave(self, gesture, hand):
        """✋ Open hand waved side to side clears this slide's drawing."""

        now = time.monotonic()

        if gesture != "Open Hand" or not self.gestures_enabled:
            self.wave_detector.idle(now)
            if self.wave_detector.progress() == 0:
                self._show_wave_progress(0)
            return False

        done = self.wave_detector.track(hand.landmark[9].x, now)

        if not done:
            self._show_wave_progress(self.wave_detector.progress())
            return False

        self.end_annotation_draw()
        self.slide_ink.clear_slide()
        self._set_action("Cleared this slide (wave)")

        # Show the full red circle briefly, then back to normal.
        self._show_wave_progress(1.0)
        QTimer.singleShot(600, lambda: self._show_wave_progress(0))

        return True

    def _show_wave_progress(self, fraction):

        if self.annotation_toolbar is not None:
            self.annotation_toolbar.set_clear_progress(fraction)

    # ============================================================
    # GESTURE PROCESSING
    #
    # Only a real, whole hand gets here (see process_frame). Gestures
    # that do something once (thumbs up / down, rock sign, three
    # fingers, pinky) must be held still for STEADY_SECONDS first
    # (thumbs and pinky: HOLD_SECONDS), so the in-between shapes of a
    # moving hand never trigger anything.
    # ============================================================

    STEADY_SECONDS = 0.3
    HOLD_SECONDS = 0.5

    # The hand moving more than this (image width) restarts the hold.
    STEADY_MOVE = 0.05

    ONE_SHOT_GESTURES = (
        "Rock & Roll",
        "Thumbs Up",
        "Thumbs Down",
        "Three Fingers",
        "Pinky",
    )

    LONG_HOLD_GESTURES = ("Thumbs Up", "Thumbs Down", "Pinky")

    def _detect(self, hand):
        """Gesture name for this frame.

        Pointing left or right with the index finger is just the
        pointer here: slides change by swiping the open hand (see
        _check_swipe).
        """

        gesture = self.gesture_detector.detect(hand)

        if gesture in ("Point Right", "Point Left"):
            return "Index Finger"

        lm = hand.landmark

        # 👍 / 👎 only from a real fist: an open palm with the thumb up
        # (or a hand turned towards the camera) is not full screen.
        if gesture in ("Thumbs Up", "Thumbs Down") and not self._closed_fist(lm):
            gesture = "None"

        # 🤙 Pinky only: close the drawing.
        if gesture == "None" and self._pinky_only(lm):
            gesture = "Pinky"

        return gesture

    @staticmethod
    def _reach_ratio(lm, tip, pip):
        """How far a fingertip reaches from the wrist, against its
        middle joint (straight ~1.4-1.6, curled ~1 or less)."""

        def reach(i):
            return ((lm[i].x - lm[0].x) ** 2 + (lm[i].y - lm[0].y) ** 2) ** 0.5

        return reach(tip) / max(reach(pip), 1e-6)

    FINGERS = ((8, 6), (12, 10), (16, 14), (20, 18))

    def _closed_fist(self, lm):
        """All four fingers curled into the palm (the thumb may be out)."""

        if SwipeDetector.hand_open(lm):
            return False

        palm = max(
            ((lm[9].x - lm[0].x) ** 2 + (lm[9].y - lm[0].y) ** 2) ** 0.5,
            1e-6,
        )
        cx, cy = SwipeDetector.palm_centre(lm)

        for tip, pip in self.FINGERS:

            # Not even loosely stretched...
            if self._reach_ratio(lm, tip, pip) > 1.1:
                return False

            # ...and folded in next to the knuckles, not pointing at
            # the camera (which also looks short).
            if ((lm[tip].x - cx) ** 2 + (lm[tip].y - cy) ** 2) ** 0.5 > palm * 0.75:
                return False

        return True

    def _pinky_only(self, lm):
        """Pinky straight, the other three fingers curled."""

        if self._reach_ratio(lm, 20, 18) < 1.3:
            return False

        return all(
            self._reach_ratio(lm, tip, pip) < 1.1
            for tip, pip in self.FINGERS[:3]
        )

    # ============================================================
    # SWIPE DEBOUNCE
    #
    # After a swipe the hand comes back to the middle, and that return
    # stroke is a quick sideways movement the other way: it used to
    # count as a second swipe (next, then previous straight away).
    # Now, after a swipe:
    #   * no swipe at all counts for SWIPE_DEBOUNCE seconds;
    #   * the hand's way back is the "return stroke": a swipe the other
    #     way while it is under way is swallowed. It is over once the
    #     hand has come back half a swipe, has been away for a moment,
    #     or after SWIPE_RETURN_SECONDS;
    #   * after the return stroke, SWIPE_REARM seconds of quiet.
    # ============================================================

    SWIPE_DEBOUNCE = 0.8
    SWIPE_RETURN_SECONDS = 2.0
    SWIPE_REARM = 0.35

    # A hand missing this long ends the return stroke (it was lowered).
    SWIPE_HAND_GONE = 0.5

    def _swipe_need(self, hand):
        """Travel (image width) the swipe detector asks of this hand."""

        detector = self.swipe_detector
        lm = hand.landmark
        palm = ((lm[9].x - lm[0].x) ** 2 + (lm[9].y - lm[0].y) ** 2) ** 0.5

        return min(
            detector.MAX_DISTANCE,
            max(detector.MIN_DISTANCE, palm * detector.SWIPE_PALMS),
        )

    def _follow_return_stroke(self, hand, now):
        """Watch the hand coming back after a swipe (every frame)."""

        back = self._swipe_return

        if back is None:
            return

        if hand is None:
            if now - back["seen"] > self.SWIPE_HAND_GONE:
                self._swipe_return = None
            return

        back["seen"] = now

        if now > back["until"]:
            self._swipe_return = None
            return

        x, _y = SwipeDetector.palm_centre(hand.landmark)
        sign = back["sign"]

        if (x - back["extreme"]) * sign > 0:
            # Still carrying on the swipe's way.
            back["extreme"] = x
        elif (back["extreme"] - x) * sign >= back["need"] * 0.5:
            # Back far enough: the return stroke is over.
            self._swipe_return = None
            self._swipe_locked_until = max(
                self._swipe_locked_until, now + self.SWIPE_REARM
            )

    def _debounce_swipe(self, swipe, hand, now):
        """True when this swipe should change the slide."""

        if now < self._swipe_locked_until:
            return False

        back = self._swipe_return
        sign = 1 if swipe == "Swipe Right" else -1

        if back is not None and sign != back["sign"]:
            # This is the hand coming back, not a new swipe.
            self._swipe_return = None
            self._swipe_locked_until = now + self.SWIPE_REARM
            return False

        x = 0.5
        need = self.swipe_detector.MIN_DISTANCE

        if hand is not None:
            x, _y = SwipeDetector.palm_centre(hand.landmark)
            need = self._swipe_need(hand)

        self._swipe_locked_until = now + self.SWIPE_DEBOUNCE
        self._swipe_return = {
            "sign": sign,
            "extreme": x,
            "need": need,
            "until": now + self.SWIPE_RETURN_SECONDS,
            "seen": now,
        }

        return True

    def _check_swipe(self, hand):
        """Next / previous slide when the open hand is swiped.

        Runs on every frame (see process_frame), whatever
        GestureDetector made of the hand: mid-swipe the hand blurs and
        turns, which reads as other gestures. `hand` is None on frames
        without a usable hand.
        """

        if not self.gestures_enabled or not self._page_opened:
            self.swipe_detector.reset()
            self._swipe_return = None
            return False

        # While drawing, an open hand waved back and forth clears the
        # slide: a swipe only counts once the hand has not turned back.
        self.swipe_detector.confirm_delay = (
            self.SWIPE_CONFIRM_SECONDS if self.annotation_mode else 0.0
        )

        now = time.monotonic()

        self._follow_return_stroke(hand, now)

        try:
            swipe = self.swipe_detector.track(hand, now)
        except Exception:
            return False

        if swipe is None or not self._debounce_swipe(swipe, hand, now):
            return False

        if swipe == "Swipe Right":
            self.next_pdf_preview()
            action = "Next Slide (swipe right)"
        else:
            self.previous_pdf_preview()
            action = "Previous Slide (swipe left)"

        self._last_swipe = (swipe, now)

        self._set_action(action)
        self.presentationAction.emit(action)
        self.detected_label.setText(swipe)
        self.update_status()

        # Keep "Next / Previous Slide" readable for a moment.
        self._swipe_shown_until = now + 1.2

        return True

    # Seconds a swipe waits while drawing (see _check_swipe).
    SWIPE_CONFIRM_SECONDS = 0.25

    # The open hand moving sideways faster than this (image widths per
    # second, over the last SWIPE_SPEED_SPAN seconds) is a swipe: the
    # annotation cursor stays still until SWIPE_FREEZE seconds after.
    SWIPE_SPEED = 0.6
    SWIPE_SPEED_SPAN = 0.15
    SWIPE_FREEZE = 0.3

    def _hand_swiping(self):
        """True while the open hand is swiping (or just swiped)."""

        now = time.monotonic()

        if now < self._swipe_locked_until or now < self._cursor_frozen_until:
            return True

        detector = self.swipe_detector

        if detector._pending:
            return True

        recent = [
            (t, x) for t, x, _y in detector.trail
            if now - t <= self.SWIPE_SPEED_SPAN
        ]

        if len(recent) >= 2:
            dt = recent[-1][0] - recent[0][0]
            if dt > 0 and abs(recent[-1][1] - recent[0][1]) / dt > self.SWIPE_SPEED:
                self._cursor_frozen_until = now + self.SWIPE_FREEZE
                return True

        return False

    # How long the arrow shows on the camera picture after a swipe.
    SWIPE_FLASH_SECONDS = 0.7

    # ============================================================
    # FINGERTIP TRAILS
    #
    # Each fingertip leaves a short fading trail on the camera
    # picture, so a swipe (and any hand movement) can be seen.
    # ============================================================

    # Landmark and colour (RGB) of each fingertip: thumb to pinky.
    TRAIL_TIPS = (
        (4, (255, 120, 90)),
        (8, (80, 170, 255)),
        (12, (90, 220, 140)),
        (16, (250, 200, 70)),
        (20, (200, 120, 255)),
    )

    TRAIL_SECONDS = 0.5

    def _update_tip_trails(self, hand, now):

        trails = self._tip_trails

        for tip, _colour in self.TRAIL_TIPS:

            trail = trails.setdefault(tip, deque())

            if hand is not None:
                point = hand.landmark[tip]
                trail.append((now, point.x, point.y))

            while trail and now - trail[0][0] > self.TRAIL_SECONDS:
                trail.popleft()

    def _draw_swipe_trail(self, frame):
        """Draw every fingertip's trail, and an arrow after a swipe, on
        the camera picture (RGB, mirrored like the trails)."""

        now = time.monotonic()

        height, width = frame.shape[:2]

        last = self._last_swipe
        flashing = (
            last is not None
            and now - last[1] < self.SWIPE_FLASH_SECONDS
        )

        for tip, colour in self.TRAIL_TIPS:

            points = [
                (t, int(x * width), int(y * height))
                for t, x, y in self._tip_trails.get(tip, ())
            ]

            # Older parts are thinner and darker, so the direction shows.
            for (t0, x0, y0), (_t1, x1, y1) in zip(points, points[1:]):
                fade = max(0.0, 1.0 - (now - t0) / self.TRAIL_SECONDS)
                shade = tuple(int(c * (0.3 + 0.7 * fade)) for c in colour)
                cv2.line(
                    frame, (x0, y0), (x1, y1), shade,
                    max(1, int(8 * fade)), cv2.LINE_AA,
                )

            if points:
                _t, x, y = points[-1]
                cv2.circle(frame, (x, y), 6, colour, -1, cv2.LINE_AA)
                cv2.circle(frame, (x, y), 8, (255, 255, 255), 1, cv2.LINE_AA)

        if not flashing:
            return

        colour = (60, 220, 120)
        right = last[0] == "Swipe Right"
        cy = height // 2
        start = int(width * (0.3 if right else 0.7))
        end = int(width * (0.7 if right else 0.3))

        cv2.arrowedLine(
            frame, (start, cy), (end, cy), colour, 10,
            cv2.LINE_AA, tipLength=0.3,
        )

        label = "NEXT" if right else "PREVIOUS"
        (text_w, _text_h), _base = cv2.getTextSize(
            label, cv2.FONT_HERSHEY_SIMPLEX, 1.0, 3
        )
        cv2.putText(
            frame, label, ((width - text_w) // 2, cy - 30),
            cv2.FONT_HERSHEY_SIMPLEX, 1.0, colour, 3, cv2.LINE_AA,
        )

    def _handle_hand(self, gesture, hand):
        """One camera frame with a real, whole hand (swipes are checked
        before this, in process_frame)."""

        if self._check_wave(gesture, hand):
            self.detected_gesture = "Wave"
            return

        now = time.monotonic()
        centre = SwipeDetector.palm_centre(hand.landmark)

        if gesture != getattr(self, "_pose", None):
            self._pose = gesture
            self._pose_since = now
            self._pose_anchor = centre

        # A hand still moving (or just after a swipe) restarts the hold.
        anchor = self._pose_anchor
        moved = ((centre[0] - anchor[0]) ** 2 + (centre[1] - anchor[1]) ** 2) ** 0.5

        if moved > self.STEADY_MOVE or now < self._swipe_locked_until:
            self._pose_since = now
            self._pose_anchor = centre

        hold = (
            self.HOLD_SECONDS
            if gesture in self.LONG_HOLD_GESTURES
            else self.STEADY_SECONDS
        )

        # One-shot gestures wait until they have been held still.
        if (
            gesture in self.ONE_SHOT_GESTURES
            and now - self._pose_since < hold
        ):
            self.end_annotation_draw()
            return

        self.process_gesture(gesture, hand)

    def _handle_no_hand(self):
        """No hand (or not a whole one) in this frame: do nothing, and
        let go of anything a hand was doing."""

        self.end_annotation_draw()

        self._pose = None
        self.detected_gesture = "None"

        # The swipe in progress is kept: a frame or two without the
        # hand is normal mid-swipe, and SwipeDetector ends it after a
        # real gap.

        if getattr(self, "wave_detector", None) is not None:
            if self.wave_detector.progress() > 0:
                self._reset_clear_hold()
            self.wave_detector.reset()

    def process_gesture(self, gesture, hand=None):

        if gesture in ("None", "Unknown"):
            self.end_annotation_draw()
            return

        # Only act once the page has been opened at least once.
        if not self._page_opened:
            return

        # 🤘 Rock sign starts the presentation (or resumes it) while
        # Presentation Control is open with slides loaded.
        if (
            gesture == "Rock & Roll"
            and not self.gestures_enabled
            and gesture != self.detected_gesture
        ):
            self.detected_gesture = gesture

            if self._has_slides() and self.isVisible():
                self.manual_start()
                self._set_action("Start Presentation")
                self.presentationAction.emit("Start Presentation")

            return

        if not self.gestures_enabled:
            self.detected_label.setText(
                gesture + " (gestures off)"
            )
            return

        # ☝️ Index finger: the pointer (a quick sideways swipe of the
        # open hand changes slides, see _check_swipe). Works while drawing too.
        if gesture == "Index Finger":
            self.end_annotation_draw()
            self.move_cursor_from_hand(hand)
            self.detected_gesture = gesture
            if time.monotonic() > getattr(self, "_swipe_shown_until", 0.0):
                self._set_action("Pointer (swipe to change slides)")
            return

        # ====================================================
        # ONE-SHOT GESTURES (same with drawing on or off)
        # ====================================================

        if gesture in self.ONE_SHOT_GESTURES:

            self.end_annotation_draw()

            # Fire once per hold.
            if gesture == self.detected_gesture:
                return

            self.detected_gesture = gesture

            if gesture in ("Thumbs Up", "Rock & Roll"):
                # 👍 Full screen (🤘 while presenting: back to it).
                self.enter_fullscreen()
                action = "Full Screen"

            elif gesture == "Thumbs Down":
                # 👎 Exit full screen.
                self.exit_fullscreen()
                action = "Exit Full Screen"

            elif gesture == "Pinky":
                # 🤙 Close the drawing: paused, so every slide keeps its
                # drawing and 🖖 three fingers opens it again.
                if not self.annotation_mode:
                    return
                self.set_annotation_mode(False)
                action = "Annotation closed"

            else:
                # 🖖 Turn drawing on. It is paused or stopped from the
                # Annotation page, so a hand passing through this shape
                # can never hide the floating circle.
                if not self.annotation_mode:
                    self.set_annotation_mode(True)
                action = "Annotation on"

            self._set_action(action)
            self.presentationAction.emit(action)
            self.update_status()
            return

        # ====================================================
        # DRAWING GESTURES
        # ====================================================

        if self.annotation_mode:

            # ✌️ Index + Middle together -> DRAW (hold left button + move).
            # Over the floating circle this drags it instead.
            if gesture == "Index + Middle":
                self.move_draw_cursor(hand)
                # Press only once the fingers have settled together,
                # so closing them does not leave a flick at the start.
                if time.monotonic() - self._pose_since >= self.DRAW_START_SECONDS:
                    self.begin_annotation_draw()
                self._set_action("Drawing")
                self.detected_gesture = gesture
                return

            # Any other gesture stops drawing.
            self.end_annotation_draw()

            # ✋ Open Hand -> move cursor over the floating circle / tools.
            # Not while the hand is swiping: the cursor stays put.
            if gesture == "Open Hand":
                if not self._hand_swiping():
                    self.move_cursor_from_hand(hand)
                    self._set_action("Cursor (select tool)")
                self.detected_gesture = gesture
                return

            # ✊ Fist -> click (opens the circle's tools, picks a tool).
            if gesture == "Fist":
                if gesture != self.detected_gesture:
                    self._safe_click()
                    self._set_action("Click")
                self.detected_gesture = gesture
                return

        self.end_annotation_draw()
        self.detected_gesture = gesture

    # ============================================================
    # ANNOTATION HELPERS
    # ============================================================

    def _set_action(self, text):

        self.current_action = text

        if hasattr(self, "status_action"):
            self.status_action.setText(text)


    def set_annotation_mode(self, enabled):
        """Three Fingers gesture: start drawing, or pause it.

        Pausing keeps every slide's drawing, the same as Pause on the
        Annotation page.
        """

        if enabled:
            self.start_live_annotation()
        else:
            self.pause_live_annotation()

    def begin_annotation_draw(self):

        if self._annot_drawing:
            return

        try:
            pyautogui.mouseDown(_pause=False)
        except Exception:
            pass

        self._annot_drawing = True

    def end_annotation_draw(self):

        if not self._annot_drawing:
            return

        try:
            pyautogui.mouseUp(_pause=False)
        except Exception:
            pass

        self._annot_drawing = False

    def _safe_click(self):

        try:
            pyautogui.click(_pause=False)
        except Exception:
            pass

    # ============================================================
    # INDEX FINGER CURSOR
    #
    # The fingertip is mapped from a box in the middle of the camera
    # picture (POINTER_BOX) to the whole screen, so every edge can be
    # reached without the hand leaving the picture. A One Euro filter
    # smooths it: steady when the finger is still, no lag when it
    # moves fast (the old fixed smoothing trailed behind the finger).
    # ============================================================

    # left, top, right, bottom of the box, as fractions of the picture.
    # A smaller box means a smaller hand movement crosses the screen.
    POINTER_BOX = (0.22, 0.15, 0.78, 0.60)

    # The pointer not used for this long starts afresh at the finger,
    # instead of sliding over from where it was last.
    POINTER_GAP = 0.3

    # ============================================================
    # ANNOTATION PEN
    #
    # Drawing needs a steadier, slower cursor than pointing:
    #   * it follows the point between the index and middle tips,
    #     which shakes less than one fingertip;
    #   * it moves DRAW_GAIN of the hand's movement, starting from
    #     where the cursor already is (no jump when drawing starts);
    #   * tremor smaller than DRAW_STEADY px is ignored (the pen hangs
    #     on a short string behind the hand), then a One Euro filter
    #     rounds off what is left.
    # ============================================================

    DRAW_GAIN = 0.55
    DRAW_STEADY = 4.0
    DRAW_START_SECONDS = 0.1

    def move_draw_cursor(self, hand):

        if hand is None or pyautogui is None:
            return

        try:

            lm = hand.landmark
            now = time.monotonic()

            x = (lm[8].x + lm[12].x) / 2.0
            y = (lm[8].y + lm[12].y) / 2.0

            left, top, right, bottom = self.POINTER_BOX
            target = (
                (x - left) / (right - left) * (self._screen_w - 1),
                (y - top) / (bottom - top) * (self._screen_h - 1),
            )

            if self._draw_hand is None or now - self._draw_last > self.POINTER_GAP:

                # Start from the cursor, wherever it is.
                cx, cy = pyautogui.position()
                self._draw_pen = (float(cx), float(cy))
                self._draw_string = self._draw_pen
                self._draw_filter_x.reset()
                self._draw_filter_y.reset()

            else:

                px, py = self._draw_pen
                hx, hy = self._draw_hand

                self._draw_pen = (
                    min(self._screen_w - 1, max(0.0, px + (target[0] - hx) * self.DRAW_GAIN)),
                    min(self._screen_h - 1, max(0.0, py + (target[1] - hy) * self.DRAW_GAIN)),
                )

            self._draw_hand = target
            self._draw_last = now

            # The string end only moves once the pen pulls it taut.
            px, py = self._draw_pen
            ex, ey = self._draw_string
            distance = ((px - ex) ** 2 + (py - ey) ** 2) ** 0.5

            if distance > self.DRAW_STEADY:
                pull = 1.0 - self.DRAW_STEADY / distance
                self._draw_string = (ex + (px - ex) * pull, ey + (py - ey) * pull)

            sx = self._draw_filter_x(self._draw_string[0], now)
            sy = self._draw_filter_y(self._draw_string[1], now)

            native_move(int(round(sx)), int(round(sy)))

            # The pointer carries on from here, not from where it was.
            self._cursor_last = 0.0

        except Exception:
            pass

    def move_cursor_from_hand(self, hand):

        if (
            hand is None
            or pyautogui is None
            or self.mp_hands is None
        ):
            return

        try:

            tip = hand.landmark[
                self.mp_hands.HandLandmark.INDEX_FINGER_TIP
            ]

            now = time.monotonic()

            if now - self._cursor_last > self.POINTER_GAP:
                self._cursor_filter_x.reset()
                self._cursor_filter_y.reset()

            self._cursor_last = now
            self._draw_hand = None

            # The camera feed is mirrored, so x already matches the
            # user's on-screen left/right.
            left, top, right, bottom = self.POINTER_BOX
            fx = min(1.0, max(0.0, (tip.x - left) / (right - left)))
            fy = min(1.0, max(0.0, (tip.y - top) / (bottom - top)))

            target_x = fx * (self._screen_w - 1)
            target_y = fy * (self._screen_h - 1)

            smooth_x = self._cursor_filter_x(target_x, now)
            smooth_y = self._cursor_filter_y(target_y, now)

            smooth_x = max(0, min(self._screen_w - 1, int(round(smooth_x))))
            smooth_y = max(0, min(self._screen_h - 1, int(round(smooth_y))))

            # Instant move: pyautogui.moveTo() sleeps 0.1 s after every
            # call, which held up each camera frame (the lag).
            native_move(smooth_x, smooth_y)

        except Exception:
            pass

    # ============================================================
    # PDF NEXT
    # ============================================================

    def next_pdf_preview(self):

        # Works for both PDF and PowerPoint (PPTX) presentations.
        if not self.pdf_document and not getattr(
            self,
            "ppt_slides",
            None
        ):
            return

        if (
            self.current_page
            < self.total_pages - 1
        ):

            self.current_page += 1

            self.render_pdf_page()

    # ============================================================
    # PDF PREVIOUS
    # ============================================================

    def previous_pdf_preview(self):

        # Works for both PDF and PowerPoint (PPTX) presentations.
        if not self.pdf_document and not getattr(
            self,
            "ppt_slides",
            None
        ):
            return

        if self.current_page > 0:

            self.current_page -= 1

            self.render_pdf_page()

    # ============================================================
    # MANUAL CONTROLS
    # ============================================================

    def manual_start(self):

        # A presentation must be loaded before it can start.
        if not self.pdf_document and not getattr(
            self,
            "ppt_slides",
            None
        ):
            QMessageBox.information(
                self,
                "No Presentation",
                "Upload a PDF or PowerPoint file first, "
                "then press Start Presentation."
            )
            return

        resuming = self.presentation_paused

        # Enable gesture control ONLY here (explicit user action).
        self.gestures_enabled = True
        self.presentation_paused = False

        self.presentation_controller.start_presentation()

        # Resuming shows the slide it was paused on (current_page is
        # kept), with that slide's drawing.
        self.enter_fullscreen()

        self.current_action = (
            "Presentation Resumed" if resuming else "Presentation Started"
        )

        self.update_status()

    def manual_pause(self):

        if not self.gestures_enabled:
            return

        # Keep the slide; gestures go off until Resume.
        self.gestures_enabled = False
        self.presentation_paused = True

        self.end_annotation_draw()
        self.exit_fullscreen()

        self.current_action = (
            f"Paused on slide {self.current_page + 1}"
        )

        self.update_status()

    def manual_stop(self):

        # Turn OFF hand-gesture control until Start is pressed again.
        self.gestures_enabled = False
        self.presentation_paused = False

        self.end_annotation_draw()

        self.presentation_controller.end_presentation()

        self.exit_fullscreen()

        self.current_action = (
            "Presentation Stopped"
        )

        self.update_status()

    def manual_next(self):

        self.presentation_controller.next_slide()

        self.next_pdf_preview()

        if self.fullscreen_window is not None:
            self.refresh_fullscreen()

        self.current_action = (
            "Next Slide"
        )

        self.update_status()

    def manual_previous(self):

        self.presentation_controller.previous_slide()

        self.previous_pdf_preview()

        if self.fullscreen_window is not None:
            self.refresh_fullscreen()

        self.current_action = (
            "Previous Slide"
        )

        self.update_status()

    # ============================================================
    # STATUS
    # ============================================================

    def update_status(self):

        self._update_counter()

        self.status_action.setText(
            self.current_action
        )

        self._refresh_controls()

    # ============================================================
    # FULLSCREEN VIEWER
    # ============================================================

    def _current_slide_image(self):
        """Return the current slide as a QImage (PDF or PPTX)."""

        if getattr(self, "ppt_slides", None):
            if -1 < self.current_page < len(self.ppt_slides):
                return self.ppt_slides[self.current_page]
            return None

        if self.pdf_document:

            page = self.pdf_document.load_page(
                self.current_page
            )

            pix = page.get_pixmap(
                matrix=fitz.Matrix(1.4, 1.4),
                alpha=False
            )

            return QImage(
                pix.samples,
                pix.width,
                pix.height,
                pix.stride,
                QImage.Format_RGB888
            )

        return None

    def enter_fullscreen(self):
        """Open (or refresh) the in-app fullscreen presentation view."""

        # If no file is loaded yet, silently do nothing (do not popup during
        # startup/splash or when the user is on another screen).
        if not self.pdf_document and not getattr(
            self,
            "ppt_slides",
            None
        ):
            return

        if self.fullscreen_window is None:
            self.fullscreen_window = FullscreenViewer(
                self,
                self
            )

        self.fullscreen_window.showFullScreen()

        self.refresh_fullscreen()

        # The annotation toolbar follows the slides into full screen.
        self._dock_annotation_bar()

    def refresh_fullscreen(self):
        """Push the current slide into the fullscreen viewer."""

        if self.fullscreen_window is not None:
            image = self._current_slide_image()

            if image is not None:
                self.fullscreen_window.set_slide(
                    image,
                    self.current_page + 1,
                    self.total_pages,
                )

    def exit_fullscreen(self):
        """Close the fullscreen view."""

        if self.fullscreen_window is not None:
            # closeEvent() hands the annotation toolbar back to the
            # preview before the viewer goes away.
            self.fullscreen_window.close()
            self.fullscreen_window = None

    # ============================================================
    # CLEANUP
    # ============================================================

    def close_camera(self):

        # Stops the camera worker first: it uses self.hands.
        self.release_camera()

        if self.hands:
            self.hands.close()

        if self.pdf_document:

            self.pdf_document.close()
            self.pdf_document = None

    def closeEvent(self, event):

        self.close_camera()

        if self.fullscreen_window is not None:
            self.fullscreen_window.close()
            self.fullscreen_window = None

        event.accept()


# ============================================================
# FULLSCREEN VIEWER
# ============================================================

class FullscreenViewer(QDialog):
    """Black fullscreen window that displays the current presentation slide.

    Created as a frameless, always-on-top, independent window (no parent) so
    it can still show fullscreen even if the main GestureBoard window is
    minimized or the user is working in another application.

    The slide is shown with its drawing from the page's SlideInk, and
    the floating annotation toolbar moves in here while it is open.
    """

    def __init__(self, parent=None, page=None):
        super().__init__()
        self.page = page

        self.setWindowTitle("Presentation")
        self.setWindowFlags(
            Qt.FramelessWindowHint
            | Qt.Window
            | Qt.WindowStaysOnTopHint
        )
        self.setStyleSheet("""
            QDialog {
                background: #000000;
            }
            QLabel {
                background: transparent;
            }
        """)

        layout = QVBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(0)

        self.slide_view = SlideView(
            page.slide_ink,
            background="#000000",
            empty_text="Loading slide…",
        )
        layout.addWidget(self.slide_view, 1)

        info = QLabel("")
        info.setObjectName("infoLabel")
        info.setAlignment(Qt.AlignCenter)
        info.setStyleSheet("""
            color: #888888;
            font-size: 12px;
            padding: 6px;
        """)
        self.info_label = info
        layout.addWidget(info)

    def set_slide(self, image, page_number, total_pages):
        """Display a slide image fullscreen."""

        self.slide_view.set_slide(QPixmap.fromImage(image))

        self.info_label.setText(
            f"{page_number} / {total_pages}"
        )

    def keyPressEvent(self, event):
        key = event.key()

        # Allow ESC / Q to exit fullscreen too.
        if key in (Qt.Key_Escape, Qt.Key_Q):
            self.close()

        # Arrow keys and presenter clickers change slides.
        elif key in (Qt.Key_Right, Qt.Key_PageDown, Qt.Key_Space):
            self.page.manual_next()
        elif key in (Qt.Key_Left, Qt.Key_PageUp):
            self.page.manual_previous()

        else:
            super().keyPressEvent(event)

    def closeEvent(self, event):
        # Clear the page's reference so Full Screen can be re-opened cleanly
        # (e.g. after the user closes the window directly), and hand the
        # annotation toolbar back to the preview so it is not deleted
        # along with this window.
        if self.page is not None:
            self.page.fullscreen_window = None
            self.page._dock_annotation_bar()
        super().closeEvent(event)
